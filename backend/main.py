import csv
import io
import logging
import traceback
from typing import Optional

from fastapi import FastAPI, UploadFile, File, Form, Depends, HTTPException, status
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy.orm import Session

from . import database, models
from .KBankPDF import parse_kbank_pdf, clean_amount
from openpyxl import load_workbook

logging.basicConfig(level=logging.INFO)

models.Base.metadata.create_all(bind=database.engine)

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


def parse_internal_ledger(file_bytes: bytes, filename: str) -> list[dict]:
    lower_name = filename.lower()

    if lower_name.endswith(".pdf"):
        return [{"note": "PDF ledger file รับเข้ามาแล้ว แต่ยังไม่ได้ parse รายการภายใน", "filename": filename}]

    if lower_name.endswith(".csv"):
        text = file_bytes.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        rows = []
        for row in reader:
            rows.append({k.strip(): v for k, v in row.items()})
        return rows

    if lower_name.endswith((".xls", ".xlsx")):
        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        sheet = workbook.active
        rows = []
        headers = []
        for row in sheet.iter_rows(values_only=True):
            if not any(cell is not None and str(cell).strip() for cell in row):
                continue
            if not headers:
                headers = [str(cell).strip() if cell is not None else f"col{idx+1}" for idx, cell in enumerate(row)]
                continue
            row_data = {}
            for idx, cell in enumerate(row):
                key = headers[idx] if idx < len(headers) else f"col{idx+1}"
                row_data[key] = "" if cell is None else str(cell)
            rows.append(row_data)
        return rows

    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail="รองรับเฉพาะไฟล์ PDF, CSV, XLSX หรือ XLS สำหรับไฟล์บัญชีภายใน"
    )

@app.post("/api/upload/statement")
async def upload_statement(
    file: UploadFile = File(...),
    password: Optional[str] = Form(None),
    db: Session = Depends(database.get_db)
):
    # 1. ตรวจสอบประเภทไฟล์
    if not file.filename.endswith(".pdf"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="รองรับเฉพาะไฟล์ PDF เท่านั้น"
        )

    try:
        # 2. อ่านไฟล์เป็น Bytes
        file_bytes = await file.read()

        # 3. ส่งไปอ่านข้อมูลผ่าน Parser
        try:
            raw_transactions = parse_kbank_pdf(file_bytes, password=password)
        except Exception as e:
            logging.exception("PDF parse error in upload_statement")
            traceback.print_exc()
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"ไม่สามารถอ่านไฟล์ Bank PDF ได้ — อาจเป็นไฟล์สแกน, เข้ารหัส หรือไฟล์ที่ไม่รองรับ: {str(e)}"
            )

        if not raw_transactions:
            return {"status": "warning", "message": "ไม่พบรายการธุรกรรมในไฟล์", "inserted_count": 0}

        # 4. บันทึก Log การอัปโหลด
        upload_log = models.UploadLog(
            filename=file.filename,
            file_type="bank_statement"
        )
        db.add(upload_log)
        db.commit()
        db.refresh(upload_log)

        # 5. แปลงข้อมูลแล้วเพิ่มลงตาราง StatementEntry
        db_entries = []
        for txn in raw_transactions:
            # รวม รายการ + ช่องทาง + รายละเอียด เข้าด้วยกันเป็น description
            desc_components = [txn["item"], txn["channel"], txn["detail"]]
            full_description = " | ".join([c for c in desc_components if c])

            entry = models.StatementEntry(
                date=txn["date"],
                description=full_description,
                amount=clean_amount(txn["amount"]),
                balance=clean_amount(txn["balance"]),
                upload_id=upload_log.id
            )
            db_entries.append(entry)

        db.bulk_save_objects(db_entries)
        db.commit()

        return {
            "status": "success",
            "upload_id": upload_log.id,
            "filename": file.filename,
            "inserted_count": len(db_entries)
        }

    except Exception as e:
        db.rollback()
        logging.exception("Error processing upload_statement")
        traceback.print_exc()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"เกิดข้อผิดพลาดในการประมวลผลไฟล์: {str(e)}"
        )


@app.post("/api/upload/files")
async def upload_files(
    bank_statement: UploadFile = File(...),
    internal_ledger: UploadFile = File(...),
    password: Optional[str] = Form(None),
    db: Session = Depends(database.get_db)
):
    if not bank_statement.filename.lower().endswith(".pdf"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="ไฟล์ bank_statement ต้องเป็น PDF เท่านั้น"
        )

    try:
        bank_bytes = await bank_statement.read()
        ledger_bytes = await internal_ledger.read()

        try:
            raw_transactions = parse_kbank_pdf(bank_bytes, password=password)
        except Exception as e:
            logging.exception("PDF parse error in upload_files")
            traceback.print_exc()
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"ไม่สามารถอ่านไฟล์ Bank PDF ได้ — อาจเป็นไฟล์สแกน, เข้ารหัส หรือไฟล์ที่ไม่รองรับ: {str(e)}"
            )

        bank_upload_log = models.UploadLog(
            filename=bank_statement.filename,
            file_type="bank_statement"
        )
        db.add(bank_upload_log)
        db.commit()
        db.refresh(bank_upload_log)

        bank_entries = []
        for txn in raw_transactions:
            desc_components = [txn["item"], txn["channel"], txn["detail"]]
            full_description = " | ".join([c for c in desc_components if c])
            entry = models.StatementEntry(
                date=txn["date"],
                description=full_description,
                amount=clean_amount(txn["amount"]),
                balance=clean_amount(txn["balance"]),
                upload_id=bank_upload_log.id
            )
            bank_entries.append(entry)

        db.bulk_save_objects(bank_entries)
        db.commit()

        ledger_rows = parse_internal_ledger(ledger_bytes, internal_ledger.filename)

        ledger_upload_log = models.UploadLog(
            filename=internal_ledger.filename,
            file_type="internal_ledger"
        )
        db.add(ledger_upload_log)
        db.commit()
        db.refresh(ledger_upload_log)

        ledger_preview = ledger_rows[:20] if isinstance(ledger_rows, list) else []

        return {
            "status": "success",
            "message": "ประมวลผลไฟล์ทั้งสองเสร็จเรียบร้อยแล้ว",
            "bank": {
                "upload_id": bank_upload_log.id,
                "filename": bank_statement.filename,
                "inserted_count": len(bank_entries)
            },
            "ledger": {
                "upload_id": ledger_upload_log.id,
                "filename": internal_ledger.filename,
                "rows_count": len(ledger_rows) if isinstance(ledger_rows, list) else 0,
                "preview": ledger_preview
            }
        }

    except HTTPException:
        raise
    except Exception as e:
        db.rollback()
        logging.exception("Error processing upload_files")
        traceback.print_exc()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"เกิดข้อผิดพลาดในการประมวลผลไฟล์: {str(e)}"
        )


@app.post("/api/upload/echo")
async def upload_echo(
    bank_statement: UploadFile = File(...),
    internal_ledger: UploadFile = File(...)
):
    """Lightweight echo endpoint for testing uploads (does not parse files)."""
    bank_bytes = await bank_statement.read()
    ledger_bytes = await internal_ledger.read()

    return {
        "status": "ok",
        "bank": {
            "filename": bank_statement.filename,
            "content_type": bank_statement.content_type,
            "size": len(bank_bytes)
        },
        "ledger": {
            "filename": internal_ledger.filename,
            "content_type": internal_ledger.content_type,
            "size": len(ledger_bytes)
        }
    }