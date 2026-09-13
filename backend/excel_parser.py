import io
from openpyxl import load_workbook
from fastapi import HTTPException

def parse_bank_excel(file_bytes: bytes) -> list[dict]:
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        headers = []
        header_idx = 0
        for idx, row in enumerate(rows[:10]):
            row_str = " ".join([str(c) for c in row if c is not None])
            if any(k in row_str for k in ["วันที่", "Date", "รายการ", "จำนวนเงิน", "Amount"]):
                headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(row)]
                header_idx = idx
                break

        if not headers:
            headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[0])]

        transactions = []
        for row in rows[header_idx + 1:]:
            if not any(c is not None for c in row):
                continue
            row_dict = {headers[i]: ("" if cell is None else str(cell).strip()) for i, cell in enumerate(row) if i < len(headers)}

            date_val = next((v for k, v in row_dict.items() if any(x in k.lower() for x in ["วัน", "date"])), "")
            time_val = next((v for k, v in row_dict.items() if any(x in k.lower() for x in ["เวลา", "time"])), "")
            item_val = next((v for k, v in row_dict.items() if any(x in k.lower() for x in ["รายการ", "desc"])), "")
            amount_val = next((v for k, v in row_dict.items() if any(x in k.lower() for x in ["จำนวน", "amount"])), "0.0")
            balance_val = next((v for k, v in row_dict.items() if any(x in k.lower() for x in ["คงเหลือ", "balance"])), "0.0")

            if date_val or amount_val:
                transactions.append({
                    "date": date_val,
                    "time": time_val,
                    "item": item_val,
                    "amount": amount_val,
                    "balance": balance_val,
                    "channel": row_dict.get("ช่องทาง", ""),
                    "detail": row_dict.get("รายละเอียด", "")
                })
        return transactions
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"อ่านไฟล์ Excel Statement ไม่สำเร็จ: {str(e)}")


def parse_internal_ledger_smart(file_bytes: bytes, filename: str) -> dict:
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))

    if not rows:
        return {"rows": [], "inserted_count": 0, "warning": "ไม่พบข้อมูลในไฟล์"}

    headers = []
    header_idx = 0
    for idx, row in enumerate(rows[:15]):
        row_str = " ".join([str(c) for c in row if c is not None])
        if any(k in row_str for k in ["วันที่", "Date", "จำนวนเงิน", "ยอดเงิน", "Price", "Amount"]):
            headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(row)]
            header_idx = idx
            break

    if not headers:
        headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[0])]

    found_fields = {"date": False, "amount": False, "time_or_id": False}
    for h in headers:
        hl = h.lower()
        if any(k in hl for k in ["วัน", "date"]): found_fields["date"] = True
        if any(k in hl for k in ["จำนวน", "ยอด", "price", "amount"]): found_fields["amount"] = True
        if any(k in hl for k in ["เวลา", "time", "เลขที่", "order", "id", "ref", "ออเดอร์"]): found_fields["time_or_id"] = True

    parsed_rows = []
    for row in rows[header_idx + 1:]:
        if not any(c is not None for c in row): continue
        row_dict = {headers[i]: ("" if cell is None else str(cell).strip()) for i, cell in enumerate(row) if i < len(headers)}
        row_str_all = " ".join(row_dict.values()).lower()

        # ตัดรายการเงินสดออก
        if "เงินสด" in row_str_all or "cash" in row_str_all:
            continue
        parsed_rows.append(row_dict)

    missing_fields, present_fields = [], []
    if found_fields["date"]: present_fields.append("วันที่")
    else: missing_fields.append("วันที่")
    if found_fields["amount"]: present_fields.append("จำนวนเงิน")
    else: missing_fields.append("จำนวนเงิน")
    if found_fields["time_or_id"]: present_fields.append("เวลา/เลขที่รายการ")
    else: missing_fields.append("เวลา/เลขที่รายการ")

    warning_msg = None
    if missing_fields:
        warning_msg = f"ไฟล์ของคุณพบฟีลด์: [{', '.join(present_fields)}] แต่ขาดฟีลด์สำคัญ: [{', '.join(missing_fields)}] ซึ่งอาจทำให้ผลการกระทบยอดคลาดเคลื่อนได้"

    return {
        "rows": parsed_rows,
        "inserted_count": len(parsed_rows),
        "warning": warning_msg,
        "found_fields": present_fields
    }